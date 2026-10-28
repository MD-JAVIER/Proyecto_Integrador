"""
generador_docx.py
-----------------
Llena la plantilla de Word de la institución (recursos/plantillas/informe_miembros.docx)
SIN cambiar su formato: se abre el mismo documento y solo se reemplaza el texto de los espacios
en blanco, conservando letra, colores, tablas y márgenes.

El informe de ofrendas se arma desde la MISMA plantilla (mismo título, tablas y estilo), cambiando
los textos y agregando las filas que hagan falta.
"""
import copy
import re
from pathlib import Path

from docx import Document
from docx.shared import RGBColor
from docx.table import Table
from docx.text.paragraph import Paragraph

PLANTILLA = Path(__file__).resolve().parents[2] / "recursos" / "plantillas" / "informe_miembros.docx"
NEGRO = RGBColor(0, 0, 0)


# ======================= utilidades (conservan el formato) =======================
def _poner_texto(parrafo, texto, color_negro=True):
    """Reemplaza el texto de un párrafo usando el formato del fragmento que tenía el texto
    (así se conservan negrita, tamaño y color, por ejemplo el texto blanco de los encabezados)."""
    if parrafo.runs:
        principal = max(parrafo.runs, key=lambda r: len(r.text))
        principal.text = texto
        for r in parrafo.runs:
            if r._r is not principal._r:          # se comparan los elementos, no los objetos envoltorio
                r.text = ""
    else:
        principal = parrafo.add_run(texto)
    if color_negro:
        principal.font.color.rgb = NEGRO
        principal.font.underline = False


def _poner_celda(celda, texto, color_negro=True):
    _poner_texto(celda.paragraphs[0], texto, color_negro)
    for extra in celda.paragraphs[1:]:         # por si la celda tenía más de un párrafo
        extra._element.getparent().remove(extra._element)


def _bloques(doc):
    """El cuerpo del documento en orden: párrafos y tablas."""
    for el in doc.element.body.iterchildren():
        tag = el.tag.split('}')[1]
        if tag == "p":
            yield Paragraph(el, doc)
        elif tag == "tbl":
            yield Table(el, doc)


def _partes(doc):
    """Separa la plantilla en: párrafos de texto, tablas y párrafos con rayas (observaciones)."""
    parrafos = [b for b in _bloques(doc) if isinstance(b, Paragraph)]
    tablas = [b for b in _bloques(doc) if isinstance(b, Table)]
    titulo, iglesia, periodo = parrafos[0], parrafos[1], parrafos[2]
    encabezados = [p for p in parrafos if p.style.name == "Heading 1"]
    rayas = [p for p in parrafos if p.text.strip() and set(p.text.strip()) == {"_"}]
    nota = parrafos[-1]
    return titulo, iglesia, periodo, encabezados, tablas, rayas, nota


def _llenar_cabecera(iglesia, periodo, p_iglesia, p_periodo, texto_periodo, fecha):
    nombre = iglesia.strip()
    _poner_texto(p_iglesia, "Nombre de la iglesia o congregación: " + (nombre if nombre else "_" * 34), False)
    _poner_texto(p_periodo,
                 f"Período del informe: {texto_periodo}    Fecha: {fecha.day:02d} / {fecha.month:02d} / {fecha.year}",
                 False)


def _llenar_observaciones(rayas, observaciones):
    """El texto escrito va en la primera línea; las otras líneas quedan para escribir a mano."""
    if observaciones.strip():
        _poner_texto(rayas[0], observaciones.strip(), False)


def _quitar_filas(tabla, posiciones):
    """Elimina filas de una tabla (se pasan las posiciones de la plantilla original)."""
    filas = list(tabla.rows)
    for i in sorted(posiciones, reverse=True):
        tabla._tbl.remove(filas[i]._tr)


def _quitar_columna(tabla, indice):
    """Quita una columna entera y reparte su ancho entre las que quedan."""
    from docx.oxml.ns import qn
    grid = tabla._tbl.tblGrid
    cols = grid.findall(qn("w:gridCol"))
    ancho = int(cols[indice].get(qn("w:w")) or 0)
    grid.remove(cols[indice])
    resto = grid.findall(qn("w:gridCol"))
    extra = ancho // max(len(resto), 1)
    for c in resto:
        c.set(qn("w:w"), str(int(c.get(qn("w:w")) or 0) + extra))
    for fila in tabla.rows:
        tcs = fila._tr.findall(qn("w:tc"))
        fila._tr.remove(tcs[indice])
        for tc in fila._tr.findall(qn("w:tc")):
            tcPr = tc.tcPr
            if tcPr is not None and tcPr.find(qn("w:tcW")) is not None:
                w = tcPr.find(qn("w:tcW"))
                w.set(qn("w:w"), str(int(w.get(qn("w:w")) or 0) + extra))
                w.set(qn("w:type"), "dxa")


def _quitar_seccion(doc, encabezado_inicio, encabezado_fin):
    """Quita del documento todo lo que hay desde un título (incluido) hasta el siguiente (sin incluir):
    el título, su tabla y los espacios que la separan."""
    cuerpo = doc.element.body
    quitar, dentro = [], False
    for el in list(cuerpo.iterchildren()):
        if el is encabezado_inicio._p:
            dentro = True
        elif el is encabezado_fin._p:
            break
        if dentro:
            quitar.append(el)
    for el in quitar:
        cuerpo.remove(el)


def _agregar_filas(tabla, cantidad, fila_modelo=1, antes_de=-1):
    """Copia una fila de la tabla (con su formato) `cantidad` veces, antes de la fila indicada."""
    destino = tabla.rows[antes_de]._tr
    for _ in range(cantidad):
        destino.addprevious(copy.deepcopy(tabla.rows[fila_modelo]._tr))


# ======================= INFORME GENERAL DE MIEMBROS =======================
def generar_informe_miembros(ruta, contenido):
    """contenido: diccionario armado por controlador_informes.datos_informe_miembros()."""
    doc = Document(str(PLANTILLA))
    titulo, p_iglesia, p_periodo, _enc, tablas, rayas, nota = _partes(doc)

    _llenar_cabecera(contenido["iglesia"], None, p_iglesia, p_periodo, contenido["periodo"], contenido["fecha"])

    # La institución solo necesita: Caballeros, Damas, Jóvenes y TOTAL; y solo los nuevos miembros.
    _quitar_columna(tablas[0], 3)           # sin "% de actividad"
    _quitar_filas(tablas[1], (4, 5))        # Niños (si corresponde) y Otras categorías
    _quitar_filas(tablas[2], (2, 3))        # Miembros retirados y Cambios de estado

    # 1. Resumen de membresía
    for celda, valor in zip(tablas[0].rows[1].cells, contenido["resumen"]):
        _poner_celda(celda, valor)

    # 2. Clasificación de miembros (las filas ya están en este orden en la plantilla)
    for fila, (cantidad, obs) in zip(tablas[1].rows[1:], contenido["clasificacion"]):
        _poner_celda(fila.cells[1], cantidad)
        _poner_celda(fila.cells[2], obs)

    # 3. Movimientos durante el período
    for fila, (cantidad, detalle) in zip(tablas[2].rows[1:], contenido["movimientos"]):
        _poner_celda(fila.cells[1], cantidad)
        _poner_celda(fila.cells[2], detalle)

    # 4. Observaciones generales
    _llenar_observaciones(rayas, contenido["observaciones"])
    _poner_texto(nota, contenido["nota"], False)

    doc.save(str(ruta))


# ======================= INFORME GENERAL DE OFRENDAS =======================
def generar_informe_ofrendas(ruta, contenido):
    """Mismo formato de la institución, aplicado a las ofrendas."""
    doc = Document(str(PLANTILLA))
    titulo, p_iglesia, p_periodo, encabezados, tablas, rayas, nota = _partes(doc)

    _poner_texto(titulo, "INFORME GENERAL DE OFRENDAS", False)
    _llenar_cabecera(contenido["iglesia"], None, p_iglesia, p_periodo, contenido["periodo"], contenido["fecha"])

    # 1. Resumen de ofrendas
    _poner_texto(encabezados[0], "1. Resumen de ofrendas", False)
    for celda, texto in zip(tablas[0].rows[0].cells,
                            ("Total del año", "Meses registrados", "Promedio mensual", "Mejor mes")):
        _poner_celda(celda, texto, False)
    for celda, valor in zip(tablas[0].rows[1].cells, contenido["resumen"]):
        _poner_celda(celda, valor)

    # 2. Detalle por mes: 12 filas + TOTAL (la plantilla trae 5 + TOTAL)
    _poner_texto(encabezados[1], "2. Detalle por mes", False)
    detalle = tablas[1]
    _agregar_filas(detalle, 12 - 5)
    for celda, texto in zip(detalle.rows[0].cells, ("Mes", "Monto", "Observaciones")):
        _poner_celda(celda, texto, False)
    for fila, (mes, monto, obs) in zip(detalle.rows[1:13], contenido["meses"]):
        _poner_celda(fila.cells[0], mes)
        _poner_celda(fila.cells[1], monto)
        _poner_celda(fila.cells[2], obs)
    total, obs_total = contenido["total"]
    _poner_celda(detalle.rows[-1].cells[1], total)
    _poner_celda(detalle.rows[-1].cells[2], obs_total)

    # La institución no usa la comparación con el año anterior: se quita la sección 3 completa
    # y "Observaciones" pasa a ser la sección 3.
    _quitar_seccion(doc, encabezados[2], encabezados[3])
    _poner_texto(encabezados[3], "3. Observaciones generales", False)

    # Observaciones generales
    _llenar_observaciones(rayas, contenido["observaciones"])
    _poner_texto(nota, contenido["nota"], False)

    doc.save(str(ruta))


# ======================= INFORME DE MIEMBROS DE UN MINISTERIO (reducido) =======================
def generar_informe_ministerio(ruta, contenido):
    """Versión reducida para el líder: resumen de SU ministerio, sus nuevos miembros y observaciones.
    Sin la clasificación por categorías (todos son del mismo ministerio)."""
    doc = Document(str(PLANTILLA))
    titulo, p_iglesia, p_periodo, encabezados, tablas, rayas, nota = _partes(doc)

    _poner_texto(titulo, f"INFORME DE MIEMBROS – {contenido['ministerio'].upper()}", False)
    _llenar_cabecera(contenido["iglesia"], None, p_iglesia, p_periodo, contenido["periodo"], contenido["fecha"])

    _quitar_columna(tablas[0], 3)                            # sin "% de actividad"
    _quitar_filas(tablas[2], (2, 3))                         # solo "Nuevos miembros"
    _quitar_seccion(doc, encabezados[1], encabezados[2])     # sin clasificación de categorías

    # 1. Resumen de membresía del ministerio
    _poner_texto(encabezados[0], f"1. Resumen de membresía – {contenido['ministerio']}", False)
    for celda, valor in zip(tablas[0].rows[1].cells, contenido["resumen"]):
        _poner_celda(celda, valor)

    # 2. Nuevos miembros
    _poner_texto(encabezados[2], "2. Nuevos miembros durante el período", False)
    cantidad, detalle = contenido["movimientos"][0]
    _poner_celda(tablas[2].rows[1].cells[1], cantidad)
    _poner_celda(tablas[2].rows[1].cells[2], detalle)

    # 3. Observaciones generales
    _poner_texto(encabezados[3], "3. Observaciones generales", False)
    _llenar_observaciones(rayas, contenido["observaciones"])
    _poner_texto(nota, contenido["nota"], False)

    doc.save(str(ruta))
